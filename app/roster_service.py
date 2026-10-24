"""Keep a tenant-scoped copy of the latest imported roster in sync with student edits."""
from __future__ import annotations

import csv
import io
import json
import logging
import os
import shutil
import tempfile
from pathlib import Path

logger = logging.getLogger("smart_campus.roster")
ROSTER_METADATA = ".student-roster.json"


def _roster_path(data_dir: Path, fmt: str) -> Path:
    return data_dir / f"student-roster.{fmt}"


def _safe_cell(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _write_xls_as_xlsx(source: Path, destination: Path) -> None:
    import xlrd
    from openpyxl import Workbook

    workbook = xlrd.open_workbook(str(source), on_demand=True)
    try:
        sheet = workbook.sheet_by_index(0)
        output = Workbook()
        target = output.active
        target.title = "Students"
        for row in range(sheet.nrows):
            values = []
            for col in range(sheet.ncols):
                cell = sheet.cell(row, col)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode).isoformat())
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    values.append(None)
                else:
                    values.append(cell.value)
            target.append(values)
        output.save(destination)
        output.close()
    finally:
        workbook.release_resources()


def retain_roster(source: Path, fmt: str, filename: str, mapping: dict[str, int],
                  data_dir: Path) -> None:
    """Save an imported CSV/XLSX roster (converting legacy XLS to XLSX) for later updates."""
    data_dir.mkdir(parents=True, exist_ok=True)
    saved_fmt = "xlsx" if fmt == "xls" else fmt
    if saved_fmt not in {"csv", "xlsx"}:
        raise ValueError(f"Unsupported roster format: {fmt}")
    destination = _roster_path(data_dir, saved_fmt)
    fd, staged_file = tempfile.mkstemp(dir=data_dir, suffix=f".{saved_fmt}.tmp")
    os.close(fd)
    metadata_path = data_dir / ROSTER_METADATA
    fd, staged_metadata = tempfile.mkstemp(dir=data_dir, suffix=".json.tmp")
    os.close(fd)
    previous_roster = data_dir / "student-roster.csv" if saved_fmt == "xlsx" else data_dir / "student-roster.xlsx"
    old_destination = destination.read_bytes() if destination.exists() else None
    old_metadata = metadata_path.read_bytes() if metadata_path.exists() else None
    try:
        if fmt == "xls":
            _write_xls_as_xlsx(source, Path(staged_file))
        else:
            shutil.copyfile(source, staged_file)
        metadata = {
            "format": saved_fmt,
            "filename": Path(filename).stem + (".xlsx" if fmt == "xls" else Path(filename).suffix.lower()),
            "mapping": mapping,
        }
        Path(staged_metadata).write_text(json.dumps(metadata), encoding="utf-8")
        os.replace(staged_file, destination)
        os.replace(staged_metadata, metadata_path)
        previous_roster.unlink(missing_ok=True)
    except Exception:
        try:
            if old_destination is None:
                destination.unlink(missing_ok=True)
            else:
                destination.write_bytes(old_destination)
            if old_metadata is None:
                metadata_path.unlink(missing_ok=True)
            else:
                metadata_path.write_bytes(old_metadata)
        except OSError:
            logger.exception("Could not roll back a retained student roster")
        for path in (staged_file, staged_metadata):
            try:
                os.unlink(path)
            except OSError:
                pass
        logger.exception("Could not retain the imported student roster")
        raise


def roster_download(data_dir: Path) -> tuple[Path, str]:
    metadata_path = data_dir / ROSTER_METADATA
    if not metadata_path.is_file():
        raise FileNotFoundError("No imported student spreadsheet is saved for this institution yet.")
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    fmt = metadata.get("format")
    if fmt not in {"csv", "xlsx"}:
        raise ValueError("The saved roster metadata is invalid.")
    path = _roster_path(data_dir, fmt)
    if not path.is_file():
        raise FileNotFoundError("The saved roster file is missing.")
    filename = metadata.get("filename")
    if not isinstance(filename, str) or Path(filename).name != filename or Path(filename).suffix.lower() != f".{fmt}":
        raise ValueError("The saved roster filename is invalid.")
    return path, filename


def stage_roster_change(data_dir: Path, student_id: str, updates: dict,
                        create: bool) -> tuple[str, str] | None:
    """Stage a corresponding append/update in the retained spreadsheet, if one exists."""
    metadata_path = data_dir / ROSTER_METADATA
    if not metadata_path.is_file():
        return None

    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    fmt = metadata.get("format")
    mapping = metadata.get("mapping")
    if fmt not in {"csv", "xlsx"} or not isinstance(mapping, dict):
        raise ValueError("The saved roster metadata is invalid.")
    student_id_col = mapping.get("student_id")
    if not isinstance(student_id_col, int) or student_id_col < 0:
        raise ValueError("The saved roster does not map a student ID column.")
    path = _roster_path(data_dir, fmt)
    if not path.is_file():
        raise FileNotFoundError("The saved roster file is missing.")

    if fmt == "csv":
        raw = path.read_bytes()
        encoding = "utf-8-sig" if raw.startswith(b"\xef\xbb\xbf") else "utf-8"
        if encoding == "utf-8":
            try:
                raw.decode(encoding)
            except UnicodeDecodeError:
                encoding = "cp1252"
        text = raw.decode(encoding)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
        if not rows:
            raise ValueError("The saved roster has no header row.")
        header = rows[0]
        for row in rows[1:]:
            row.extend([""] * (len(header) - len(row)))
        for index in mapping.values():
            if not isinstance(index, int) or index < 0 or index >= len(header):
                raise ValueError("The saved roster column mapping is invalid.")
        match_index = next((i for i, row in enumerate(rows[1:], 1)
                            if len(row) > student_id_col and row[student_id_col].strip().upper() == student_id), None)
        if create:
            if match_index is not None:
                raise ValueError(f"Student '{student_id}' already exists in the saved roster.")
            target = [""] * len(header)
            rows.append(target)
            match_index = len(rows) - 1
        elif match_index is None:
            raise ValueError(f"Student '{student_id}' was not found in the saved roster.")
        row = rows[match_index]
        for key, index in mapping.items():
            if key in updates:
                row[index] = "" if updates[key] is None else _safe_cell(str(updates[key]))
        fd, staged = tempfile.mkstemp(dir=data_dir, suffix=".csv.tmp")
        try:
            with os.fdopen(fd, "w", newline="", encoding=encoding) as handle:
                csv.writer(handle, dialect).writerows(rows)
        except Exception:
            os.unlink(staged)
            raise
        return staged, str(path)

    from openpyxl import load_workbook

    workbook = load_workbook(path)
    staged: str | None = None
    try:
        worksheet = workbook.active
        for index in mapping.values():
            if not isinstance(index, int) or index < 0 or index >= worksheet.max_column:
                raise ValueError("The saved roster column mapping is invalid.")
        match_row = next((row for row in range(2, worksheet.max_row + 1)
                          if str(worksheet.cell(row, student_id_col + 1).value or "").strip().upper() == student_id), None)
        if create:
            if match_row is not None:
                raise ValueError(f"Student '{student_id}' already exists in the saved roster.")
            match_row = worksheet.max_row + 1
        elif match_row is None:
            raise ValueError(f"Student '{student_id}' was not found in the saved roster.")
        for key, index in mapping.items():
            if key in updates:
                worksheet.cell(match_row, index + 1, "" if updates[key] is None else _safe_cell(updates[key]))
        fd, staged = tempfile.mkstemp(dir=data_dir, suffix=".xlsx.tmp")
        os.close(fd)
        workbook.save(staged)
        return staged, str(path)
    except Exception:
        if staged:
            try:
                os.unlink(staged)
            except OSError:
                pass
        raise
    finally:
        workbook.close()
