"""
Bulk student import (additive feature module).

This module never changes existing backend behaviour. It reuses the same CSV working files,
schemas (config.DATASET_SCHEMAS) and validation rules as the single-student create/update API,
and reloads the shared DataService once after a successful import so analytics refresh.

Pipeline
  upload   -> file is stored in a private temp folder; headers + first rows are previewed
  validate -> background job: parse in bounded batches, validate every row server-side
  import   -> background job: re-validate (client results are never trusted), then commit

Persistence model: the working data is a set of CSV files, so the commit stages every affected CSV
and swaps them in one step (with in-memory rollback if any swap fails). Nothing is written until
validation of the whole file has finished, which is why cancellation is safe until the commit step.
"""
from __future__ import annotations

import csv
import difflib
import io
import logging
import os
import re
import secrets
import tempfile
import threading
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

import pandas as pd

from .config import DATASET_FILES, DATASET_SCHEMAS
from .data_service import DataService, DataValidationError, PersistenceError
from .roster_service import retain_roster

logger = logging.getLogger("smart_campus.import")

# ----------------------------------------------------------------------------- limits
MAX_UPLOAD_BYTES = int(os.getenv("IMPORT_MAX_UPLOAD_MB", "10")) * 1024 * 1024
MAX_ROWS = int(os.getenv("IMPORT_MAX_ROWS", "20000"))
MAX_COLUMNS = 200
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024   # zip-bomb guard for .xlsx
BATCH_SIZE = 500                              # rows processed between progress / cancel checks
PREVIEW_ROWS = 100
MAX_STORED_ISSUES = 50_000
UPLOAD_TTL_SECONDS = 2 * 60 * 60
ALLOWED_EXTENSIONS = (".xlsx", ".xls", ".csv")

MODES = ("insert_only", "update_existing", "insert_or_update")
MODE_LABELS = {
    "insert_only": "Insert only — add new students; skip IDs that already exist",
    "update_existing": "Update existing — change existing students; skip IDs that are not found",
    "insert_or_update": "Insert or update — add new students and update existing ones",
}
ID_PATTERN = re.compile(r"^STU[0-9]{8}$")
EMAIL_HINT = "Enter a valid email address."


class ImportError_(Exception):
    """User-facing import problem (mapped to an HTTP error by the router)."""

    def __init__(self, message: str, status: int = 422):
        super().__init__(message)
        self.message, self.status = message, status


# ----------------------------------------------------------------------------- field catalogue
@dataclass(frozen=True)
class FieldSpec:
    key: str
    label: str
    kind: str                       # text | int | float | category
    group: str
    required_for_insert: bool = False
    minimum: float | None = None
    maximum: float | None = None
    options: tuple[str, ...] = ()
    aliases: tuple[str, ...] = ()
    example: str = ""
    max_length: int | None = None


_INT_FIELDS = {
    "year_of_study", "admission_year", "backlogs", "subjects_below_attendance_threshold",
    "logins_last_30_days", "events_attended", "hackathons_participated",
    "certifications_count", "projects_count",
}
_LABELS = {
    "student_id": "Student ID", "student_name": "Student name", "department": "Department",
    "year_of_study": "Year of study", "section": "Section", "admission_year": "Admission year",
    "student_email": "Student email", "cgpa": "CGPA", "backlogs": "Backlogs",
    "internal_score_pct": "Internal score %", "semester_score_pct": "Semester score %",
    "overall_attendance_pct": "Overall attendance %",
    "subjects_below_attendance_threshold": "Subjects below attendance threshold",
    "logins_last_30_days": "LMS logins (30 days)", "assignment_completion_pct": "Assignment completion %",
    "avg_weekly_learning_hours": "Weekly learning hours", "events_attended": "Events attended",
    "hackathons_participated": "Hackathons participated", "club_participation": "Club participation",
    "volunteer_hours": "Volunteer hours", "aptitude_score": "Aptitude score", "coding_score": "Coding score",
    "mock_interview_score": "Mock interview score", "resume_readiness_pct": "Resume readiness %",
    "placement_status": "Placement status", "technical_skill_score": "Technical skill score",
    "soft_skill_score": "Soft skill score", "certifications_count": "Certifications",
    "projects_count": "Projects", "student_satisfaction_1_to_5": "Student satisfaction (1–5)",
    "faculty_feedback_1_to_5": "Faculty feedback (1–5)", "support_requested": "Support requested",
    "feedback_topic": "Feedback topic",
}
_ALIASES = {
    "student_id": ("roll no", "roll number", "rollno", "reg no", "registration number", "enrollment number",
                   "enrolment no", "usn", "id", "student no", "student number"),
    "student_name": ("name", "full name", "student", "candidate name"),
    "department": ("dept", "branch", "stream", "program", "programme", "course"),
    "year_of_study": ("year", "study year", "current year", "yr"),
    "section": ("sec", "division", "class section"),
    "admission_year": ("year of admission", "joining year", "batch", "admitted year", "intake year"),
    "student_email": ("email", "email id", "mail", "email address", "e mail"),
    "cgpa": ("gpa", "cumulative gpa", "grade point"),
    "backlogs": ("arrears", "backlog count", "number of backlogs", "failed subjects"),
    "internal_score_pct": ("internal", "internal marks", "internal percentage", "internals"),
    "semester_score_pct": ("semester", "semester marks", "semester percentage", "sem score"),
    "overall_attendance_pct": ("attendance", "attendance percentage", "attendance pct", "overall attendance"),
    "assignment_completion_pct": ("assignment completion", "assignments completed"),
    "avg_weekly_learning_hours": ("weekly hours", "learning hours", "study hours"),
    "logins_last_30_days": ("lms logins", "logins", "monthly logins"),
    "support_requested": ("needs support", "support needed"),
    "placement_status": ("placement", "placed status"),
}
_EXAMPLES = {
    "student_id": "STU20260501", "student_name": "Ananya Rao", "department": "Computer Science and Engineering",
    "year_of_study": "2", "section": "A", "admission_year": "2025", "student_email": "ananya@example.edu",
    "cgpa": "7.8", "club_participation": "Active", "placement_status": "Eligible", "support_requested": "No",
    "feedback_topic": "Campus Engagement",
}
_REQUIRED_INSERT = {"student_id", "student_name", "department", "year_of_study", "section", "admission_year"}
_MAX_LENGTHS = {"student_name": 120, "department": 120, "section": 8, "student_email": 254, "feedback_topic": 120}
_CHOICE_OVERRIDES = {"year_of_study": (1, 8), "admission_year": (1990, 2100)}


def _build_catalogue() -> list[FieldSpec]:
    """Derive the importable fields from the same schema config the CSV store uses (single source of truth)."""
    specs: list[FieldSpec] = [FieldSpec(
        "student_id", _LABELS["student_id"], "text", "Identity", aliases=_ALIASES["student_id"],
        required_for_insert=True, example=_EXAMPLES["student_id"])]
    seen = {"student_id"}
    group_names = {"students": "Student profile", "academics": "Academics", "attendance": "Attendance",
                   "lms": "Learning platform", "engagement": "Engagement", "placements": "Placement",
                   "skills": "Skills", "feedback": "Feedback"}
    for dataset, schema in DATASET_SCHEMAS.items():
        if dataset == "interventions":
            continue
        for col in schema["required"]:
            if col in seen:
                continue
            seen.add(col)
            lo, hi = schema["ranges"].get(col, (None, None))
            if col in _CHOICE_OVERRIDES:
                lo, hi = _CHOICE_OVERRIDES[col]
            options = tuple(schema["categories"].get(col, ()))
            kind = "category" if options else "int" if col in _INT_FIELDS else "float" if lo is not None else "text"
            specs.append(FieldSpec(
                key=col, label=_LABELS.get(col, col.replace("_", " ").title()), kind=kind,
                group=group_names[dataset], required_for_insert=col in _REQUIRED_INSERT,
                minimum=lo, maximum=hi, options=options, aliases=_ALIASES.get(col, ()),
                example=_EXAMPLES.get(col, ""), max_length=_MAX_LENGTHS.get(col)))
    return specs


FIELDS: list[FieldSpec] = _build_catalogue()
FIELD_BY_KEY: dict[str, FieldSpec] = {f.key: f for f in FIELDS}


def field_catalogue() -> list[dict]:
    return [{"key": f.key, "label": f.label, "kind": f.kind, "group": f.group,
             "required_for_insert": f.required_for_insert, "minimum": f.minimum, "maximum": f.maximum,
             "options": list(f.options), "example": f.example} for f in FIELDS]


def template_csv() -> str:
    """Header row plus one illustrative row. Every importable field is included."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([f.key for f in FIELDS])
    sample = {"student_id": "STU20260501", "student_name": "Ananya Rao", "department": "Computer Science and Engineering",
              "year_of_study": "2", "section": "A", "admission_year": "2025",
              "student_email": "ananya.rao@example.edu", "cgpa": "7.8", "backlogs": "0",
              "internal_score_pct": "68", "semester_score_pct": "72", "overall_attendance_pct": "84.5",
              "subjects_below_attendance_threshold": "1", "logins_last_30_days": "18",
              "assignment_completion_pct": "76", "avg_weekly_learning_hours": "5.5", "events_attended": "3",
              "hackathons_participated": "1", "club_participation": "Active", "volunteer_hours": "12",
              "aptitude_score": "70", "coding_score": "64", "mock_interview_score": "62",
              "resume_readiness_pct": "58", "placement_status": "Eligible", "technical_skill_score": "66",
              "soft_skill_score": "71", "certifications_count": "2", "projects_count": "3",
              "student_satisfaction_1_to_5": "4", "faculty_feedback_1_to_5": "4", "support_requested": "No",
              "feedback_topic": "Campus Engagement"}
    writer.writerow([sample.get(f.key, "") for f in FIELDS])
    return out.getvalue()


# ----------------------------------------------------------------------------- parsing
def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return ""
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def detect_format(filename: str, head: bytes) -> str:
    """Return xlsx | xls | csv, checking that the file content matches its extension."""
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ImportError_("Unsupported file type. Upload an Excel (.xlsx, .xls) or CSV (.csv) file.", 415)
    is_zip, is_ole = head.startswith(b"PK\x03\x04"), head.startswith(b"\xd0\xcf\x11\xe0")
    if ext == ".xlsx" and not is_zip:
        raise ImportError_("This file has an .xlsx extension but is not a valid Excel workbook.", 422)
    if ext == ".xls" and not is_ole:
        raise ImportError_("This file has an .xls extension but is not a valid legacy Excel workbook.", 422)
    if ext == ".csv" and (is_zip or is_ole):
        raise ImportError_("This file has a .csv extension but looks like an Excel workbook. "
                           "Rename it to .xlsx/.xls or export it as CSV.", 422)
    if ext == ".csv" and b"\x00" in head:
        raise ImportError_("This file does not look like a text CSV file.", 422)
    return ext[1:]


def _iter_csv(path: Path) -> Iterator[list[str]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost everything
        raise ImportError_("The CSV file encoding is not supported. Save it as UTF-8.")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    for row in csv.reader(io.StringIO(text, newline=""), dialect):
        yield [c.strip() for c in row]


def _iter_xlsx(path: Path) -> Iterator[list[str]]:
    import openpyxl
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ImportError_("The workbook is too large when decompressed.", 413)
    except zipfile.BadZipFile as exc:
        raise ImportError_("The workbook is corrupt or not a valid .xlsx file.") from exc
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_("The workbook could not be opened. It may be corrupt or password-protected.") from exc
    try:
        sheet = next((ws for ws in workbook.worksheets if ws.sheet_state == "visible"), workbook.worksheets[0])
        for row in sheet.iter_rows(values_only=True):
            yield [_cell_text(c) for c in row]
    finally:
        workbook.close()


def _iter_xls(path: Path) -> Iterator[list[str]]:
    import xlrd
    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:
        raise ImportError_("The .xls workbook could not be opened. It may be corrupt or password-protected.") from exc
    try:
        sheet = book.sheet_by_index(0)
        for r in range(sheet.nrows):
            out: list[str] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        out.append(_cell_text(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                    except Exception:
                        out.append(str(cell.value))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    out.append("")
                else:
                    out.append(_cell_text(cell.value))
            yield out
    finally:
        book.release_resources()


def iter_sheet_rows(path: Path, fmt: str) -> Iterator[list[str]]:
    return {"csv": _iter_csv, "xlsx": _iter_xlsx, "xls": _iter_xls}[fmt](path)


@dataclass
class ParsedSheet:
    headers: list[str]
    rows: list[tuple[int, list[str]]]   # (spreadsheet row number, cells padded to header width)
    blank_rows: int = 0


def parse_sheet(path: Path, fmt: str, limit: int | None = None) -> ParsedSheet:
    """Read the header and data rows. Spreadsheet row numbers are 1-based; the header is the first non-blank row."""
    headers: list[str] | None = None
    rows: list[tuple[int, list[str]]] = []
    blank = 0
    for number, cells in enumerate(iter_sheet_rows(path, fmt), start=1):
        if headers is None:
            if not any(cells):
                continue
            while cells and not cells[-1]:
                cells = cells[:-1]
            if len(cells) > MAX_COLUMNS:
                raise ImportError_(f"The file has more than {MAX_COLUMNS} columns.")
            headers = [h or f"Column {i + 1}" for i, h in enumerate(cells)]
            continue
        if not any(cells):
            blank += 1
            continue
        if len(rows) >= MAX_ROWS:
            raise ImportError_(f"The file has more than {MAX_ROWS:,} data rows. Split it into smaller files.", 413)
        width = len(headers)
        rows.append((number, (cells + [""] * width)[:width]))
        if limit is not None and len(rows) >= limit:
            break
    if headers is None:
        raise ImportError_("The file is empty: no header row was found.")
    if not rows and limit is None:
        raise ImportError_("The file has a header row but no data rows.")
    return ParsedSheet(headers, rows, blank)


# ----------------------------------------------------------------------------- mapping suggestions
def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", text.lower().replace("_", " ")).strip()


def suggest_mapping(headers: list[str]) -> dict[str, dict]:
    """Return {field_key: {column, confidence, reason}} choosing each column at most once."""
    normed = [_norm(h) for h in headers]
    candidates: list[tuple[float, str, int, str]] = []
    for spec in FIELDS:
        key_norm, label_norm = _norm(spec.key), _norm(spec.label)
        alias_norms = {_norm(a) for a in spec.aliases}
        for idx, h in enumerate(normed):
            if not h:
                continue
            if h == key_norm:
                candidates.append((1.0, spec.key, idx, "Exact column name"))
            elif h == label_norm:
                candidates.append((0.98, spec.key, idx, "Matches field label"))
            elif h in alias_norms:
                candidates.append((0.92, spec.key, idx, "Known alias"))
            else:
                ratio = max(difflib.SequenceMatcher(None, h, key_norm).ratio(),
                            difflib.SequenceMatcher(None, h, label_norm).ratio())
                if ratio >= 0.82:
                    candidates.append((round(0.6 + ratio * 0.3, 2), spec.key, idx, "Similar name"))
    candidates.sort(key=lambda c: (-c[0], c[2]))
    used_cols: set[int] = set()
    out: dict[str, dict] = {}
    for conf, key, idx, reason in candidates:
        if key in out or idx in used_cols:
            continue
        out[key] = {"column": idx, "confidence": conf, "reason": reason}
        used_cols.add(idx)
    return out


def normalise_mapping(raw: dict[str, Any], column_count: int, mode: str) -> dict[str, int]:
    """Validate a client mapping {field_key: column_index}. Returns a clean mapping or raises ImportError_."""
    if mode not in MODES:
        raise ImportError_(f"Unknown import mode '{mode}'. Choose one of: {', '.join(MODES)}.")
    mapping: dict[str, int] = {}
    for key, col in (raw or {}).items():
        if col is None or col == "":
            continue
        if key not in FIELD_BY_KEY:
            raise ImportError_(f"'{key}' is not an importable student field.")
        if isinstance(col, bool) or not isinstance(col, int) or not 0 <= col < column_count:
            raise ImportError_(f"The column chosen for '{FIELD_BY_KEY[key].label}' does not exist in the file.")
        mapping[key] = col
    dupes = {c for c in mapping.values() if list(mapping.values()).count(c) > 1}
    if dupes:
        fields = [FIELD_BY_KEY[k].label for k, c in mapping.items() if c in dupes]
        raise ImportError_(f"One spreadsheet column is mapped to more than one field: {', '.join(fields)}.")
    problems: list[str] = []
    if mode != "insert_only" and "student_id" not in mapping:
        problems.append("Student ID (the unique identifier used to match existing students)")
    if mode in ("insert_only", "insert_or_update"):
        problems += [FIELD_BY_KEY[k].label for k in sorted(_REQUIRED_INSERT) if k not in mapping]
    if mode == "update_existing" and not [k for k in mapping if k != "student_id"]:
        problems.append("at least one field to update")
    if problems:
        raise ImportError_("Map the required fields before continuing: " + ", ".join(problems) + ".")
    return mapping


# ----------------------------------------------------------------------------- row validation
def _convert(spec: FieldSpec, text: str) -> tuple[Any, str | None]:
    """Convert one cell. Returns (value, error_message)."""
    if spec.kind in ("int", "float"):
        cleaned = text.replace(",", "").rstrip("%").strip() if spec.key.endswith("_pct") else text.replace(",", "")
        try:
            number = float(cleaned)
        except ValueError:
            return None, f"“{text}” is not a number."
        if number != number or number in (float("inf"), float("-inf")):
            return None, f"“{text}” is not a valid number."
        if spec.kind == "int":
            if not number.is_integer():
                return None, f"“{text}” must be a whole number."
            number = int(number)
        if spec.minimum is not None and number < spec.minimum or spec.maximum is not None and number > spec.maximum:
            return None, f"{text} is outside the allowed range {spec.minimum:g}–{spec.maximum:g}."
        return number, None
    if spec.kind == "category":
        lookup = {o.lower(): o for o in spec.options}
        found = lookup.get(text.lower())
        if found is None:
            return None, f"“{text}” is not allowed. Use one of: {', '.join(spec.options)}."
        return found, None
    if spec.key == "student_email":
        if "@" not in text or text.startswith("@") or text.endswith("@") or " " in text:
            return None, EMAIL_HINT
    if spec.max_length and len(text) > spec.max_length:
        return None, f"Text is {len(text)} characters; the limit is {spec.max_length}."
    return text, None


@dataclass
class Issue:
    row: int
    field: str | None
    column: int | None
    value: str
    code: str
    message: str

    def as_dict(self, headers: list[str]) -> dict:
        return {"row": self.row, "field": self.field, "field_label": FIELD_BY_KEY[self.field].label if self.field else None,
                "column": self.column, "source_column": headers[self.column] if self.column is not None else None,
                "value": self.value, "code": self.code, "message": self.message}


@dataclass
class RowPlan:
    row: int
    student_id: str | None
    action: str                    # insert | update
    values: dict[str, Any]


@dataclass
class ValidationResult:
    total_rows: int = 0
    plans: list[RowPlan] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)
    issues_truncated: bool = False
    invalid_rows: set[int] = field(default_factory=set)
    duplicate_rows: set[int] = field(default_factory=set)
    unchanged_rows: int = 0
    to_insert: int = 0
    to_update: int = 0
    skipped_existing: int = 0
    skipped_not_found: int = 0
    generated_ids: int = 0

    def add(self, issue: Issue) -> None:
        if len(self.issues) < MAX_STORED_ISSUES:
            self.issues.append(issue)
        else:
            self.issues_truncated = True

    def summary(self) -> dict:
        invalid, duplicates = len(self.invalid_rows), len(self.duplicate_rows)
        return {"total_rows": self.total_rows, "valid_rows": len(self.plans), "duplicate_rows": duplicates,
                "invalid_rows": invalid, "will_insert": self.to_insert, "will_update": self.to_update,
                "skipped_existing": self.skipped_existing, "skipped_not_found": self.skipped_not_found,
                "unchanged_rows": self.unchanged_rows, "generated_ids": self.generated_ids,
                "issue_count": len(self.issues), "issues_truncated": self.issues_truncated}


class Cancelled(Exception):
    pass


def validate_rows(parsed: ParsedSheet, mapping: dict[str, int], mode: str, existing_ids: set[str],
                  progress: Callable[[int, int], None] | None = None,
                  should_cancel: Callable[[], bool] | None = None) -> ValidationResult:
    """Validate every row in bounded batches. Pure function of its inputs (no I/O)."""
    result = ValidationResult(total_rows=len(parsed.rows))
    first_seen: dict[str, int] = {}
    id_col = mapping.get("student_id")
    data_fields = [(k, c) for k, c in mapping.items() if k != "student_id"]
    allow_insert, allow_update = mode in ("insert_only", "insert_or_update"), mode in ("update_existing", "insert_or_update")

    for batch_start in range(0, len(parsed.rows), BATCH_SIZE):
        if should_cancel and should_cancel():
            raise Cancelled()
        for number, cells in parsed.rows[batch_start: batch_start + BATCH_SIZE]:
            row_issues: list[Issue] = []
            sid: str | None = None
            raw_id = cells[id_col].strip() if id_col is not None else ""
            if raw_id:
                sid = raw_id.upper()
                if not ID_PATTERN.match(sid):
                    row_issues.append(Issue(number, "student_id", id_col, raw_id, "invalid_id",
                                            "Student ID must look like STU followed by 8 digits (e.g. STU20260501)."))
                    sid = None
            elif mode != "insert_only":
                row_issues.append(Issue(number, "student_id", id_col, "", "missing_id",
                                        "Student ID is required in this mode to match or create the record."))

            # Duplicate IDs inside the file: first occurrence wins, later ones are reported and skipped.
            if sid and sid in first_seen:
                result.duplicate_rows.add(number)
                result.add(Issue(number, "student_id", id_col, raw_id, "duplicate_in_file",
                                 f"Duplicate of row {first_seen[sid]} in this file. This row will be skipped."))
                continue

            exists = sid in existing_ids if sid else False
            values: dict[str, Any] = {}
            for key, col in data_fields:
                text = cells[col].strip()
                if not text:
                    continue
                value, error = _convert(FIELD_BY_KEY[key], text)
                if error:
                    row_issues.append(Issue(number, key, col, text, "invalid_value", error))
                else:
                    values[key] = value

            # Required fields when a brand-new student would be created.
            will_insert = allow_insert and not exists and not (sid is None and raw_id)
            if will_insert:
                if not raw_id and not any(i.field == "student_id" for i in row_issues):
                    row_issues.append(Issue(number, "student_id", id_col, "", "required_missing",
                                            "Student ID is required to add a new student."))
                for key in sorted(_REQUIRED_INSERT - {"student_id"}):
                    col = mapping.get(key)
                    if key not in values and not any(i.field == key for i in row_issues):
                        row_issues.append(Issue(number, key, col, "", "required_missing",
                                                f"{FIELD_BY_KEY[key].label} is required to add a new student."))

            if row_issues:
                result.invalid_rows.add(number)
                for issue in row_issues:
                    result.add(issue)
                if sid:
                    first_seen.setdefault(sid, number)
                continue

            if sid:
                first_seen[sid] = number
            if exists and mode == "insert_only":
                result.duplicate_rows.add(number)
                result.skipped_existing += 1
                result.add(Issue(number, "student_id", id_col, raw_id, "already_exists",
                                 f"{sid} already exists. Insert-only mode skips it; use Update or Insert-or-update to change it."))
                continue
            if not exists and mode == "update_existing":
                result.invalid_rows.add(number)
                result.skipped_not_found += 1
                result.add(Issue(number, "student_id", id_col, raw_id, "not_found",
                                 f"{sid} was not found. Update-existing mode cannot create students."))
                continue
            if exists:
                if not values:
                    result.unchanged_rows += 1
                result.to_update += 1
                result.plans.append(RowPlan(number, sid, "update", values))
            else:
                if sid is None:
                    result.generated_ids += 1
                result.to_insert += 1
                result.plans.append(RowPlan(number, sid, "insert", values))
        if progress:
            progress(min(batch_start + BATCH_SIZE, len(parsed.rows)), len(parsed.rows))
    return result


# ----------------------------------------------------------------------------- commit (CSV working files)
def _fmt(value: Any) -> str:
    return "" if value is None else str(value)


def commit_plans(service: DataService, plans: list[RowPlan], department_lookup: Callable[[str], str | None],
                 should_cancel: Callable[[], bool] | None = None,
                 on_point_of_no_return: Callable[[], None] | None = None) -> dict:
    """Apply validated plans to every student CSV in one staged swap, then reload analytics once."""
    inserted = updated = 0
    with service._write_lock:
        if should_cancel and should_cancel():
            raise Cancelled()
        data_dir = service.data_dir
        existing_numeric = [int(s[3:]) for s in service.students.index if str(s).startswith("STU") and str(s)[3:].isdigit()]
        next_number = max(existing_numeric, default=20260000) + 1
        used_ids = set(service.students.index)
        resolved: list[tuple[str, RowPlan]] = []
        for plan in plans:
            sid = plan.student_id
            if sid is None:
                while f"STU{next_number:08d}" in used_ids:
                    next_number += 1
                sid = f"STU{next_number:08d}"
            used_ids.add(sid)
            resolved.append((sid, plan))

        originals: dict[str, bytes | None] = {}
        staged: list[tuple[str, str]] = []
        try:
            for dataset, filename in DATASET_FILES.items():
                if dataset == "interventions":
                    continue
                path = data_dir / filename
                columns = DATASET_SCHEMAS[dataset]["required"]
                originals[str(path)] = path.read_bytes() if path.exists() else None
                frame = pd.read_csv(path, dtype=str, keep_default_na=False) if path.exists() else pd.DataFrame(columns=columns)
                for col in columns:
                    if col not in frame.columns:
                        frame[col] = ""
                frame = frame[columns].copy()
                index = {str(v).strip().upper(): i for i, v in enumerate(frame["student_id"])}
                new_rows: list[dict[str, str]] = []
                for sid, plan in resolved:
                    values = {k: v for k, v in plan.values.items() if k in columns and k != "student_id"}
                    if "department" in values:
                        values["department"] = department_lookup(str(values["department"])) or values["department"]
                    if plan.action == "insert":
                        row = {c: "" for c in columns}
                        row["student_id"] = sid
                        row.update({k: _fmt(v) for k, v in values.items()})
                        new_rows.append(row)
                    elif sid in index:
                        for k, v in values.items():
                            frame.at[index[sid], k] = _fmt(v)
                    elif dataset != "students":
                        # Existing student with no row yet in this feature file: create it so the update is not lost.
                        row = {c: "" for c in columns}
                        row["student_id"] = sid
                        row.update({k: _fmt(v) for k, v in values.items()})
                        new_rows.append(row)
                if new_rows:
                    frame = pd.concat([frame, pd.DataFrame(new_rows, columns=columns)], ignore_index=True)
                fd, tmp_name = tempfile.mkstemp(dir=str(data_dir), suffix=".tmp")
                with os.fdopen(fd, "w", newline="", encoding="utf-8") as handle:
                    frame.to_csv(handle, index=False)
                staged.append((tmp_name, str(path)))
            if should_cancel and should_cancel():
                raise Cancelled()          # last safe cancellation point: nothing has been replaced yet
            if on_point_of_no_return:
                on_point_of_no_return()    # from here on the job can no longer be cancelled
            swapped: list[str] = []
            try:
                for tmp_name, final_path in staged:
                    os.replace(tmp_name, final_path)
                    swapped.append(final_path)
                service.load()
            except Exception:
                for final_path in swapped:       # roll back files already replaced
                    original = originals.get(final_path)
                    if original is not None:
                        Path(final_path).write_bytes(original)
                try:
                    service.load()
                except Exception:               # pragma: no cover
                    logger.exception("Reload after rollback failed")
                raise
            inserted = sum(1 for _, p in resolved if p.action == "insert")
            updated = len(resolved) - inserted
        except (Cancelled, DataValidationError):
            raise
        except Exception as exc:
            logger.exception("Bulk import commit failed")
            raise PersistenceError("The import could not be saved. No changes were applied.") from exc
        finally:
            for tmp_name, _ in staged:
                try:
                    os.unlink(tmp_name)
                except OSError:
                    pass
    return {"inserted": inserted, "updated": updated}


# ----------------------------------------------------------------------------- uploads + jobs
@dataclass
class Upload:
    upload_id: str
    owner: str
    filename: str
    fmt: str
    path: Path
    size: int
    created: float
    headers: list[str]
    total_rows: int
    blank_rows: int


@dataclass
class Job:
    job_id: str
    upload_id: str
    kind: str                         # validate | import
    owner: str
    mode: str
    filename: str
    headers: list[str]
    status: str = "queued"            # queued | running | completed | failed | cancelled
    phase: str = "Queued"
    processed: int = 0
    total: int = 0
    cancel_requested: bool = False
    can_cancel: bool = True
    error: str | None = None
    summary: dict | None = None
    result: dict | None = None
    issues: list[Issue] = field(default_factory=list)
    created: float = field(default_factory=time.time)
    finished: float | None = None

    def public(self) -> dict:
        pct = int(round(self.processed / self.total * 100)) if self.total else (100 if self.status == "completed" else 0)
        return {"job_id": self.job_id, "upload_id": self.upload_id, "kind": self.kind, "mode": self.mode,
                "status": self.status, "phase": self.phase, "processed": self.processed, "total": self.total,
                "percent": min(100, pct), "can_cancel": self.can_cancel and self.status in ("queued", "running"),
                "cancel_requested": self.cancel_requested, "error": self.error, "summary": self.summary,
                "result": self.result, "filename": self.filename, "created_at": self.created,
                "finished_at": self.finished}


class ImportManager:
    """In-process registry of uploads and background jobs (single worker = one write at a time)."""

    def __init__(self, storage_dir: Path | None = None):
        base = storage_dir or Path(os.getenv("IMPORT_STORAGE_DIR", "") or Path(tempfile.gettempdir()) / "academia_imports")
        self.storage_dir = Path(base)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        self.uploads: dict[str, Upload] = {}
        self.jobs: dict[str, Job] = {}
        self._lock = threading.RLock()
        self._executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="student-import")

    # --- uploads
    def purge_expired(self) -> None:
        cutoff = time.time() - UPLOAD_TTL_SECONDS
        with self._lock:
            for uid in [u for u, up in self.uploads.items() if up.created < cutoff]:
                self.uploads.pop(uid).path.unlink(missing_ok=True)
            for jid in [j for j, job in self.jobs.items() if job.created < cutoff and job.status not in ("queued", "running")]:
                self.jobs.pop(jid)

    def save_upload(self, owner: str, filename: str, stream: Iterable[bytes]) -> Upload:
        self.purge_expired()
        safe_name = re.sub(r"[^A-Za-z0-9._ -]", "_", Path(filename or "upload").name)[:120] or "upload"
        upload_id = secrets.token_urlsafe(16)
        path = self.storage_dir / f"{upload_id}{Path(safe_name).suffix.lower()}"
        size, head = 0, b""
        try:
            with path.open("wb") as handle:
                for chunk in stream:
                    size += len(chunk)
                    if size > MAX_UPLOAD_BYTES:
                        raise ImportError_(f"The file is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.", 413)
                    if len(head) < 16:
                        head += chunk[:16]
                    handle.write(chunk)
            if size == 0:
                raise ImportError_("The file is empty.")
            fmt = detect_format(safe_name, head)
            parsed = parse_sheet(path, fmt)
        except BaseException:
            path.unlink(missing_ok=True)
            raise
        upload = Upload(upload_id, owner, safe_name, fmt, path, size, time.time(), parsed.headers,
                        len(parsed.rows), parsed.blank_rows)
        with self._lock:
            self.uploads[upload_id] = upload
        return upload

    def get_upload(self, upload_id: str, owner: str) -> Upload:
        with self._lock:
            upload = self.uploads.get(upload_id)
        if upload is None or upload.owner != owner or not upload.path.exists():
            raise ImportError_("This upload was not found or has expired. Upload the file again.", 404)
        return upload

    def discard_upload(self, upload_id: str, owner: str) -> None:
        upload = self.get_upload(upload_id, owner)
        with self._lock:
            if any(j.upload_id == upload_id and j.status in ("queued", "running") for j in self.jobs.values()):
                raise ImportError_("A job is still running for this upload.", 409)
            self.uploads.pop(upload_id, None)
        upload.path.unlink(missing_ok=True)

    # --- jobs
    def get_job(self, job_id: str, owner: str) -> Job:
        with self._lock:
            job = self.jobs.get(job_id)
        if job is None or job.owner != owner:
            raise ImportError_("Job not found.", 404)
        return job

    def cancel(self, job_id: str, owner: str) -> Job:
        job = self.get_job(job_id, owner)
        if job.status in ("queued", "running"):
            if not job.can_cancel:
                raise ImportError_("The import is already being saved and can no longer be cancelled.", 409)
            job.cancel_requested = True
        return job

    def submit(self, upload: Upload, kind: str, mapping: dict[str, int], mode: str, service: DataService) -> Job:
        with self._lock:
            if kind == "import" and any(j.kind == "import" and j.status in ("queued", "running") for j in self.jobs.values()):
                raise ImportError_("Another import is already running. Wait for it to finish.", 409)
            job = Job(secrets.token_urlsafe(12), upload.upload_id, kind, upload.owner, mode, upload.filename, upload.headers)
            self.jobs[job.job_id] = job
        self._executor.submit(self._run, job, upload, mapping, service)
        return job

    def _run(self, job: Job, upload: Upload, mapping: dict[str, int], service: DataService) -> None:
        def cancelled() -> bool:
            return job.cancel_requested

        try:
            job.status, job.phase = "running", "Reading file"
            if cancelled():
                raise Cancelled()
            parsed = parse_sheet(upload.path, upload.fmt)
            job.total = len(parsed.rows)
            job.phase = "Validating rows"

            def progress(done: int, total: int) -> None:
                job.processed, job.total = done, total

            existing = set(service.students.index)
            result = validate_rows(parsed, mapping, job.mode, existing, progress, cancelled)
            job.issues, job.summary = result.issues, result.summary()
            if job.kind == "validate":
                job.status, job.phase, job.processed = "completed", "Validation complete", job.total
                return

            if not result.plans:
                job.result = {**result.summary(), "inserted": 0, "updated": 0, "skipped": job.total,
                              "message": "No new changes were needed."}
                if (result.plans or result.skipped_existing or result.skipped_not_found
                        or result.unchanged_rows):
                    try:
                        retain_roster(upload.path, upload.fmt, upload.filename, mapping, service.data_dir)
                        job.result["roster_saved"] = True
                    except Exception:
                        logger.exception("Roster could not be retained after a no-change import")
                        job.result["roster_saved"] = False
                        job.result["message"] = (
                            "No student changes were needed, but the uploaded spreadsheet could not be saved for future updates."
                        )
                job.status, job.phase, job.processed = "completed", "Nothing to import", job.total
                return
            job.phase, job.processed, job.total = "Saving changes", 0, len(result.plans)
            job.can_cancel = True
            def lock_in() -> None:
                job.can_cancel = False
            outcome = commit_plans(service, result.plans, service.match_department, cancelled, lock_in)
            job.processed = job.total
            skipped = job.summary["total_rows"] - len(result.plans)
            job.result = {**outcome, "skipped": skipped, "total_rows": job.summary["total_rows"],
                          "invalid_rows": job.summary["invalid_rows"], "duplicate_rows": job.summary["duplicate_rows"],
                          "generated_ids": job.summary["generated_ids"], "issue_count": job.summary["issue_count"],
                          "roster_saved": True}
            try:
                retain_roster(upload.path, upload.fmt, upload.filename, mapping, service.data_dir)
            except Exception:
                logger.exception("Student import succeeded, but its source roster could not be retained")
                job.result["roster_saved"] = False
                job.result["message"] = (
                    "Students were imported, but the uploaded spreadsheet could not be saved for future updates."
                )
            job.status, job.phase = "completed", "Import complete"
        except Cancelled:
            job.status, job.phase = "cancelled", "Cancelled — no changes were saved"
            job.result = {"inserted": 0, "updated": 0} if job.kind == "import" else None
        except ImportError_ as exc:
            job.status, job.phase, job.error = "failed", "Failed", exc.message
        except (PersistenceError, DataValidationError) as exc:
            job.status, job.phase, job.error = "failed", "Failed", f"{exc} No changes were applied."
        except Exception:
            logger.exception("Import job failed")
            job.status, job.phase, job.error = "failed", "Failed", "Unexpected server error. No changes were applied."
        finally:
            job.can_cancel = False
            job.finished = time.time()


# ----------------------------------------------------------------------------- reports
def _safe_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text   # CSV/Excel formula-injection guard


def error_report_csv(job: Job) -> str:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["spreadsheet_row", "field", "source_column", "value", "error_code", "description"])
    for issue in sorted(job.issues, key=lambda i: (i.row, i.field or "")):
        d = issue.as_dict(job.headers)
        writer.writerow([d["row"], _safe_cell(d["field_label"] or ""), _safe_cell(d["source_column"] or ""),
                         _safe_cell(d["value"]), d["code"], _safe_cell(d["message"])])
    return out.getvalue()
